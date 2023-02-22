import sys
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.chart import BarChart,Reference
wb = Workbook()
ws = wb.active

date_2021 = [["Localitatea","Timisoara", "Domeniul de incadrare a ratei de criminalitate"],["Ridicat","Coeficient de "
                                                                                                      "criminalitate "
                                                                                                      "local",
                                                                                            111.15],[ "An", 2021,
                                                                                                      "Localitatea",
                                                                                                      "Timisoara",],
             ["Domeniul de incadrare a ratei de criminalitate","Ridicat","Coeficient de criminalitate local"],
             [117.19, "An", 2022]]
for row in date_2021:
    ws.append(row)

ft = Font(bold=True)
for row in ws['A1:d1']:
    for cell in row:
        cell_font = ft
chart = BarChart()
chart.type = "col"
chart.title = "Rata Criminalitatii"
chart.y_axis.title = "Valoarea "
chart.x_axis.title = "Date Analizate"
chart.legend = None


data = Reference(ws, min_col=3, min_row=2, max_row=3, max_col=4)
categories = Reference(ws, min_col=1, min_row=2, max_row=3, max_col=4)

chart.add_data(data)
chart.set_categories(categories)

ws.add_chart(chart,'E1')
wb.save("TreeData.xlsx")
pd = pd.DataFrame(date_2021)
print(pd)
print("Thank You My Friend!")
